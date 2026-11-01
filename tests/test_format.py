import io
import math
import unittest

import pandas as pd
from openpyxl import load_workbook

from src.core.nlp_competitor import description_summary, generate_description
from src.utils.format import (
    EMPTY,
    angka,
    column_kind,
    desimal,
    excel_bytes,
    parse_rupiah,
    persen,
    rupiah,
    style_frame,
    table_height,
)


class FormatTest(unittest.TestCase):
    def test_numbers(self):
        self.assertEqual(rupiah(1431776359), "Rp 1.431.776.359")
        self.assertEqual(rupiah(-5000), "−Rp 5.000")
        self.assertEqual(angka(8162), "8.162")
        self.assertEqual(angka(8161.6), "8.162")
        self.assertEqual(desimal(21.0812), "21,08")
        self.assertEqual(desimal(3.14159, 3), "3,142")
        self.assertEqual(persen(22.63), "22,6%")
        for empty in (None, float("nan"), math.inf, "abc"):
            self.assertEqual(rupiah(empty), EMPTY)
            self.assertEqual(angka(empty), EMPTY)
            self.assertEqual(desimal(empty), EMPTY)
            self.assertEqual(persen(empty), EMPTY)

    def test_parse_rupiah(self):
        self.assertEqual(parse_rupiah("Rp 250.000.000"), 250_000_000)
        self.assertEqual(parse_rupiah("250000000"), 250_000_000)
        self.assertEqual(parse_rupiah("1.500.000,75"), 1_500_000)
        self.assertEqual(parse_rupiah(""), 0)
        self.assertEqual(parse_rupiah(None), 0)

    def test_column_kind(self):
        self.assertEqual(column_kind("Sales volume"), "money")
        self.assertEqual(column_kind("Biaya iklan"), "money")
        self.assertEqual(column_kind("Kontribusi invoice"), "pct")
        self.assertEqual(column_kind("ROAS"), "dec")
        self.assertEqual(column_kind("QTY terjual"), "count")
        self.assertEqual(column_kind("Qty retur"), "count")

    def test_table_height(self):
        self.assertEqual(table_height(3), 4 * 35 + 3)
        self.assertEqual(table_height(1000), 600)
        self.assertEqual(table_height(0), 38)

    def test_style_frame_display_only(self):
        frame = pd.DataFrame({"Produk": ["A", "Total"], "Sales": [1500000.0, 1500000.0], "ROAS": [8.123, None]})
        html = style_frame(frame, bold_label="Total").to_html()
        self.assertIn("Rp 1.500.000", html)
        self.assertIn("8,12", html)
        self.assertIn(EMPTY, html)
        self.assertIn("font-weight: 700", html)
        self.assertEqual(frame["Sales"].dtype.kind, "f")

    def test_excel_numeric(self):
        frame = pd.DataFrame({"SKU": ["1"], "Sales": [1431776359.0], "ROAS": [21.0812]})
        book = load_workbook(io.BytesIO(excel_bytes(frame)))
        ws = book["Data"]
        self.assertEqual(ws["B2"].value, 1431776359)
        self.assertEqual(ws["B2"].number_format, "#,##0")
        self.assertEqual(ws["C2"].number_format, "0.00")


class DescriptionTest(unittest.TestCase):
    def test_no_capital_title_and_summary(self):
        text = generate_description(
            title="Kloset Duduk TOTO CW421J Putih",
            name="KLOSET DUDUK TOTO CW421J PUTIH",
            variation="",
            description="",
            primary=["kloset duduk", "toto"],
            supporting=["closet"],
            cancel_reasons=[],
            max_chars=3000,
        )
        self.assertFalse(text.startswith("KLOSET DUDUK TOTO"))
        summary = description_summary(text)
        self.assertLessEqual(len(summary), 60)
        self.assertTrue(summary.endswith("…") or len(summary) < 60)


if __name__ == "__main__":
    unittest.main()
