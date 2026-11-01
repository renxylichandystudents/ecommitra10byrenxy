"""Format angka Indonesia, tabel tanpa baris kosong, input rupiah, dan Excel numerik."""

from __future__ import annotations

import io
import re
from typing import Iterable

import numpy as np
import pandas as pd

EMPTY = "–"
ROW_PX = 35
MONEY_HINTS = ("sales", "omzet", "biaya", "rp", "budget", "nilai", "harga", "alokasi", "gmv", "penjualan", "modal", "target")
PCT_HINTS = ("%", "persen", "porsi", "tingkat", "kontribusi", "ctr", "acos", "konversi (%)")
DECIMAL_HINTS = ("roas", "mape", "rmse", "wape", "bias", "mae", "rasio", "efektivitas", "efektifitas", "skor", "bobot")
COUNT_HINTS = ("qty", "unit", "pesanan", "jumlah", "terjual", "stok", "klik", "dilihat", "peringkat", "no", "baris", "sku dengan")


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if np.isfinite(number) else None


def _group(number: float, decimals: int) -> str:
    text = f"{number:,.{decimals}f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def angka(value: object) -> str:
    """8162 -> "8.162"."""
    number = _number(value)
    return EMPTY if number is None else _group(round(number), 0)


def rupiah(value: object) -> str:
    """1431776359 -> "Rp 1.431.776.359"; negatif -> "−Rp 5.000"."""
    number = _number(value)
    if number is None:
        return EMPTY
    sign = "−" if number < 0 else ""
    return f"{sign}Rp {_group(abs(round(number)), 0)}"


def desimal(value: object, n: int = 2) -> str:
    """21.0812 -> "21,08"."""
    number = _number(value)
    return EMPTY if number is None else _group(number, n)


def persen(value: object, n: int = 1) -> str:
    """22.6 -> "22,6%". Nilai sudah dalam persen (bukan pecahan)."""
    number = _number(value)
    return EMPTY if number is None else f"{_group(number, n)}%"


def parse_rupiah(text: object) -> int:
    """"Rp 250.000.000" / "250000000" / "250.000.000,50" -> 250000000."""
    raw = str(text or "").strip()
    raw = raw.split(",")[0]
    digits = re.sub(r"[^0-9]", "", raw)
    return int(digits) if digits else 0


def column_kind(name: object) -> str:
    """Jenis format kolom dari namanya: money, pct, dec, atau count."""
    label = str(name).strip().lower()
    if any(hint in label for hint in PCT_HINTS):
        return "pct"
    if any(hint in label for hint in DECIMAL_HINTS):
        return "dec"
    if any(hint in label for hint in COUNT_HINTS) and not any(hint in label for hint in ("sales", "omzet", "biaya", "nilai")):
        return "count"
    if any(re.search(rf"\b{re.escape(hint)}\b", label) for hint in MONEY_HINTS):
        return "money"
    return "count"


def formatter_for(kind: str, series: pd.Series):
    if kind == "money":
        return rupiah
    if kind == "pct":
        return persen
    if kind == "dec":
        return desimal
    values = pd.to_numeric(series, errors="coerce").dropna()
    integral = values.empty or bool(np.allclose(values, np.round(values)))
    return angka if integral else desimal


def table_height(rows: int, max_height: int = 600) -> int:
    return int(min(max_height, (max(int(rows), 0) + 1) * ROW_PX + 3))


BLANK_CELLS = {"", "-", "–", "—", "nan", "none", "nat", "<na>"}


def drop_blank_rows(frame: pd.DataFrame) -> pd.DataFrame:
    """Buang baris yang semua selnya kosong (NaN/None/""/"-")."""
    if frame is None or frame.empty or not len(frame.columns):
        return frame
    candidate = pd.Series(True, index=range(len(frame)))
    for i in range(len(frame.columns)):
        rows = candidate.to_numpy().nonzero()[0]
        if not len(rows):
            return frame
        column = frame.iloc[rows, i]
        missing = column.isna().to_numpy(dtype=bool)
        blank = missing.copy()
        if not (pd.api.types.is_numeric_dtype(column) or pd.api.types.is_datetime64_any_dtype(column)):
            text = [str(value).strip().lower() for value in column.astype(object).tolist()]
            blank |= np.fromiter((value in BLANK_CELLS for value in text), dtype=bool, count=len(text))
        candidate.iloc[rows] = blank
    return frame.iloc[~candidate.to_numpy()]


def style_frame(
    frame: pd.DataFrame,
    *,
    kinds: dict[str, str] | None = None,
    skip: Iterable[str] = (),
    bold_label: str | None = None,
):
    """Styler format tampilan saja; nilai sel tetap numerik agar sort benar."""
    kinds = dict(kinds or {})
    skip = set(skip)
    formats = {}
    for column in frame.columns:
        if column in skip or not pd.api.types.is_numeric_dtype(frame[column]) or pd.api.types.is_bool_dtype(frame[column]):
            continue
        formats[column] = formatter_for(kinds.get(column) or column_kind(column), frame[column])
    styler = frame.style.format(formats, na_rep=EMPTY)
    if bold_label is not None and len(frame.columns):
        first = frame.columns[0]
        label_column = next((c for c in frame.columns if frame[c].astype(str).eq(bold_label).any()), first)
        mask = frame[label_column].astype(str).eq(bold_label)
        styler = styler.apply(lambda row: ["font-weight:700" if mask.loc[row.name] else "" for _ in row], axis=1)
    return styler


def _has_own_format(cfg: object) -> bool:
    """Kolom progress/format khusus dibiarkan; kolom yang hanya berisi help tetap diformat Styler."""
    if not isinstance(cfg, dict):
        return False
    type_config = cfg.get("type_config") or {}
    own = type_config.get("format") not in (None, "", "localized")
    return own or type_config.get("type") == "progress"


def help_columns(tips: dict[str, str], frame: pd.DataFrame, labels: dict[str, str] | None = None) -> dict:
    """column_config berisi tooltip saja (format angka tetap dari Styler)."""
    import streamlit as st

    labels = labels or {}
    return {name: st.column_config.Column(labels.get(name), help=text) for name, text in tips.items() if name in frame.columns}


def show_table(
    frame: pd.DataFrame,
    max_height: int = 600,
    *,
    kinds: dict[str, str] | None = None,
    bold_label: str | None = None,
    column_config: dict | None = None,
    styled: bool = True,
    **kwargs,
):
    """st.dataframe dengan tinggi pas isi (tanpa baris kosong) dan format angka Indonesia.

    Baris yang seluruh selnya kosong (NaN, "", "-") dibuang. Tabel kosong → "Belum ada data".
    Tinggi baris dikunci ROW_PX agar rumus tinggi tetap pas walau ukuran font browser berbeda.
    """
    import streamlit as st

    if frame is None:
        frame = pd.DataFrame()
    if "on_select" not in kwargs:
        frame = drop_blank_rows(frame)
    if frame.empty or not len(frame.columns):
        st.info("Belum ada data")
        return None
    config = dict(column_config or {})
    preformatted = [name for name, cfg in config.items() if _has_own_format(cfg)]
    data = style_frame(frame, kinds=kinds, skip=preformatted, bold_label=bold_label) if styled else frame
    return st.dataframe(
        data,
        height=table_height(len(frame), max_height),
        row_height=ROW_PX,
        hide_index=kwargs.pop("hide_index", True),
        use_container_width=kwargs.pop("use_container_width", True),
        column_config=config or None,
        **kwargs,
    )


def input_rupiah(label: str, value: int | float, key: str, *, help: str | None = None) -> int:
    """text_input tampil "250.000.000"; terima ketikan dengan/tanpa titik; kembalikan int."""
    import streamlit as st

    if key not in st.session_state:
        st.session_state[key] = angka(value)

    def _reformat() -> None:
        st.session_state[key] = angka(parse_rupiah(st.session_state.get(key)))

    text = st.text_input(f"{label} (Rp)", key=key, on_change=_reformat, help=help)
    return parse_rupiah(text)


def _excel_ready(frame: pd.DataFrame) -> pd.DataFrame:
    clean = frame.copy()
    for column in clean.columns:
        if clean[column].dtype == object and clean[column].map(lambda v: isinstance(v, (list, dict, set, tuple))).any():
            clean[column] = clean[column].astype(str)
    return clean


def excel_bytes(frames: dict[str, pd.DataFrame] | pd.DataFrame) -> bytes:
    """Excel dengan angka asli + number format '#,##0' / '0.00' (bukan teks)."""
    from openpyxl.utils import get_column_letter

    if isinstance(frames, pd.DataFrame):
        frames = {"Data": frames}
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in frames.items():
            clean = _excel_ready(frame)
            sheet = str(name)[:31] or "Data"
            clean.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            for position, column in enumerate(clean.columns, start=1):
                series = clean[column]
                if not pd.api.types.is_numeric_dtype(series) or pd.api.types.is_bool_dtype(series):
                    continue
                values = pd.to_numeric(series, errors="coerce").dropna()
                integral = values.empty or bool(np.allclose(values, np.round(values)))
                kind = column_kind(column)
                fmt = "0.00" if (kind in {"dec", "pct"} or not integral) else "#,##0"
                letter = get_column_letter(position)
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
                ws.column_dimensions[letter].width = max(12, min(40, len(str(column)) + 4))
    return buffer.getvalue()
